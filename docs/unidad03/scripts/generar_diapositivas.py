# -*- coding: utf-8 -*-
"""
Genera docs/unidad03/Unidad-3-Diapositivas-2026.pptx a partir del contenido de
docs/unidad03/04-diapositivas.md (36 diapositivas).

Requiere: pip install python-pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ----------------------------------------------------------------------------
# Paleta y constantes
# ----------------------------------------------------------------------------
NAVY = RGBColor(0x1F, 0x2A, 0x44)
NAVY_LIGHT = RGBColor(0x33, 0x41, 0x63)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
RED = RGBColor(0xDC, 0x26, 0x26)
AMBER = RGBColor(0xD9, 0x77, 0x06)
TEAL = RGBColor(0x0D, 0x94, 0x88)
SLATE = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
LIGHT_BG = RGBColor(0xF3, 0xF4, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CODE_BG = RGBColor(0x0F, 0x17, 0x2A)
CODE_TEXT = RGBColor(0xE2, 0xE8, 0xF0)
CODE_GREEN = RGBColor(0x4A, 0xDE, 0x80)
CODE_COMMENT = RGBColor(0x94, 0xA3, 0xB8)

FONT_TITLE = "Calibri"
FONT_BODY = "Calibri"
FONT_CODE = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN_X = Inches(0.55)
HEADER_H = Inches(1.05)
FOOTER_Y = Inches(7.08)

TOTAL_SLIDES = 36


def set_bg(slide, color=WHITE):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, color, line=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line:
        shp.line.color.rgb = color
        shp.line.width = Pt(0.5)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_pill(slide, x, y, text, bg_color, text_color=WHITE, font_size=11, width=None):
    w = width or Inches(0.35 + 0.085 * len(text))
    h = Inches(0.32)
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        shp.adjustments[0] = 0.5
    except Exception:
        pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = bg_color
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(font_size)
    r.font.bold = True
    r.font.color.rgb = text_color
    r.font.name = FONT_BODY
    return shp


def add_textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tb, tf


def add_header(slide, num, title, badge=None, kicker=None):
    add_rect(slide, 0, 0, SLIDE_W, HEADER_H, NAVY)
    add_rect(slide, 0, HEADER_H, SLIDE_W, Pt(3), TEAL)

    # numero de slide, circulo
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.3), Inches(0.24), Inches(0.58), Inches(0.58))
    circ.fill.solid()
    circ.fill.fore_color.rgb = TEAL
    circ.line.fill.background()
    circ.shadow.inherit = False
    ctf = circ.text_frame
    ctf.margin_left = 0; ctf.margin_right = 0; ctf.margin_top = 0; ctf.margin_bottom = 0
    cp = ctf.paragraphs[0]
    cp.alignment = PP_ALIGN.CENTER
    cr = cp.add_run()
    cr.text = str(num)
    cr.font.size = Pt(18)
    cr.font.bold = True
    cr.font.color.rgb = WHITE
    cr.font.name = FONT_TITLE

    tb, tf = add_textbox(slide, Inches(1.05), Inches(0.14), Inches(9.6), Inches(0.85), anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        pk = tf.paragraphs[0]
        rk = pk.add_run()
        rk.text = kicker.upper()
        rk.font.size = Pt(11)
        rk.font.bold = True
        rk.font.color.rgb = RGBColor(0x9C, 0xC9, 0xC5)
        rk.font.name = FONT_BODY
        p = tf.add_paragraph()
    else:
        p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = FONT_TITLE

    if badge:
        add_pill(slide, Inches(10.85), Inches(0.37), badge, AMBER, WHITE, font_size=11, width=Inches(2.1))


def add_footer(slide, num):
    add_textbox_simple(slide, Inches(0.55), FOOTER_Y, Inches(9), Inches(0.3),
                        "Unidad III — Pruebas Unitarias con Unity · Cryptbound build 0.4",
                        size=10, color=MUTED, bold=False, italic=True)
    add_textbox_simple(slide, Inches(12.3), FOOTER_Y, Inches(0.7), Inches(0.3),
                        f"{num}/{TOTAL_SLIDES}", size=10, color=MUTED, bold=False, align=PP_ALIGN.RIGHT)


def add_textbox_simple(slide, x, y, w, h, text, size=18, color=SLATE, bold=False,
                        italic=False, align=PP_ALIGN.LEFT, font=FONT_BODY, anchor=MSO_ANCHOR.TOP):
    tb, tf = add_textbox(slide, x, y, w, h, anchor=anchor)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return tb


def add_bullets(slide, x, y, w, h, items, size=18, color=SLATE, gap=8, bullet_color=TEAL):
    tb, tf = add_textbox(slide, x, y, w, h)
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        r = p.add_run()
        r.text = f"■  {item}"
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = FONT_BODY
    return tb


def add_tag(slide, x, y, label):
    add_pill(slide, x, y, f"Ejemplo: {label}", TEAL, WHITE, font_size=12,
             width=Inches(0.6 + 0.09 * len(label)))


def add_pregunta(slide, text):
    y = Inches(6.15)
    box = add_rect(slide, MARGIN_X, y, SLIDE_W - 2 * MARGIN_X, Inches(0.72), RGBColor(0xEA, 0xF2, 0xF1))
    tb, tf = add_textbox(slide, MARGIN_X + Inches(0.2), y, SLIDE_W - 2 * MARGIN_X - Inches(0.4), Inches(0.72),
                          anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = f"❓  {text}"
    r.font.size = Pt(15)
    r.font.italic = True
    r.font.bold = True
    r.font.color.rgb = NAVY
    r.font.name = FONT_BODY


def add_notes(slide, notes):
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def new_slide(prs):
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    set_bg(slide)
    return slide


# ----------------------------------------------------------------------------
# Slides de contenido
# ----------------------------------------------------------------------------

def slide_title(prs, s):
    slide = new_slide(prs)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    add_rect(slide, 0, Inches(4.35), SLIDE_W, Pt(3), TEAL)
    add_textbox_simple(slide, Inches(1), Inches(2.15), Inches(11.3), Inches(0.6),
                        "UNIDAD III", size=20, color=RGBColor(0x9C, 0xC9, 0xC5), bold=True)
    add_textbox_simple(slide, Inches(1), Inches(2.65), Inches(11.3), Inches(1.5),
                        s["title"], size=44, color=WHITE, bold=True)
    add_textbox_simple(slide, Inches(1), Inches(3.75), Inches(11.3), Inches(0.6),
                        s["subtitle"], size=20, color=RGBColor(0xD1, 0xD5, 0xDB))
    add_textbox_simple(slide, Inches(1), Inches(4.55), Inches(11.3), Inches(0.5),
                        s["footer"], size=14, color=RGBColor(0x9C, 0xC9, 0xC5), italic=True)
    # marca visual PASS/FAIL -> PASS
    add_pill(slide, Inches(1), Inches(5.25), "❌ FAIL", RED, WHITE, font_size=14, width=Inches(1.3))
    add_textbox_simple(slide, Inches(2.45), Inches(5.22), Inches(0.6), Inches(0.4), "→", size=20, color=WHITE)
    add_pill(slide, Inches(3.0), Inches(5.25), "✅ PASS", GREEN, WHITE, font_size=14, width=Inches(1.3))
    add_footer(slide, s["num"])
    add_notes(slide, s.get("notes"))


def slide_bullets(prs, s):
    slide = new_slide(prs)
    add_header(slide, s["num"], s["title"], badge=s.get("badge"), kicker=s.get("kicker"))
    y = Inches(1.45)
    if s.get("tag"):
        add_tag(slide, MARGIN_X, y, s["tag"])
        y += Inches(0.55)
    if s.get("quote"):
        add_textbox_simple(slide, MARGIN_X, y, SLIDE_W - 2 * MARGIN_X, Inches(1.1),
                            s["quote"], size=22, italic=True, color=NAVY, bold=True)
        y += Inches(1.15)
    if s.get("bullets"):
        h = Inches(6.0) - y
        add_bullets(slide, MARGIN_X, y, SLIDE_W - 2 * MARGIN_X, h, s["bullets"], size=s.get("bsize", 19))
    if s.get("pregunta"):
        add_pregunta(slide, s["pregunta"])
    add_footer(slide, s["num"])
    add_notes(slide, s.get("notes"))


def slide_flow(prs, s):
    slide = new_slide(prs)
    add_header(slide, s["num"], s["title"], badge=s.get("badge"))
    steps = s["flow"]
    n = len(steps)
    top = Inches(1.5)
    bottom = Inches(6.0) if not s.get("pregunta") else Inches(5.95)
    avail = bottom - top
    box_h = min(Inches(0.62), Emu(int(avail / n) - Inches(0.18)))
    gap = Emu(int((avail - box_h * n) / max(n - 1, 1))) if n > 1 else 0
    y = top
    box_w = Inches(9.9)
    x = (SLIDE_W - box_w) / 2
    font_size = 17 if n <= 4 else (15 if n == 5 else 13)
    for i, step in enumerate(steps):
        is_last = (i == n - 1)
        color = GREEN if (is_last and s.get("highlight_last")) else NAVY_LIGHT
        add_rect(slide, x, y, box_w, box_h, color)
        tb, tf = add_textbox(slide, x + Inches(0.25), y, box_w - Inches(0.5), box_h, anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = step
        r.font.size = Pt(font_size)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.name = FONT_BODY
        y_arrow = y + box_h
        if not is_last:
            add_textbox_simple(slide, x, y_arrow, box_w, gap, "↓", size=20, color=TEAL,
                                align=PP_ALIGN.CENTER, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        y = y_arrow + gap
    if s.get("pregunta"):
        add_pregunta(slide, s["pregunta"])
    add_footer(slide, s["num"])
    add_notes(slide, s.get("notes"))


def slide_table(prs, s):
    slide = new_slide(prs)
    add_header(slide, s["num"], s["title"])
    headers = s["headers"]
    rows = s["rows"]
    ncols = len(headers)
    nrows = len(rows) + 1
    x, y = MARGIN_X, Inches(1.5)
    w = SLIDE_W - 2 * MARGIN_X
    h = Inches(4.5)
    gtable = slide.shapes.add_table(nrows, ncols, x, y, w, h).table
    for c, htext in enumerate(headers):
        cell = gtable.cell(0, c)
        cell.text = htext
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(14)
                r.font.color.rgb = WHITE
                r.font.name = FONT_BODY
    for ri, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = gtable.cell(ri, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT_BG if ri % 2 == 0 else WHITE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if c == 1 else PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.size = Pt(13)
                    r.font.color.rgb = SLATE
                    r.font.name = FONT_BODY
    if s.get("notes_line"):
        add_textbox_simple(slide, MARGIN_X, Inches(6.15), SLIDE_W - 2 * MARGIN_X, Inches(0.7),
                            s["notes_line"], size=14, italic=True, color=MUTED)
    add_footer(slide, s["num"])
    add_notes(slide, s.get("notes"))


def add_code_block(slide, x, y, w, h, code_lines):
    add_rect(slide, x, y, w, h, CODE_BG)
    tb, tf = add_textbox(slide, x + Inches(0.25), y + Inches(0.18), w - Inches(0.5), h - Inches(0.36))
    first = True
    for line in code_lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = line if line != "" else " "
        r.font.size = Pt(14.5)
        r.font.name = FONT_CODE
        r.font.color.rgb = CODE_GREEN if line.strip().startswith("//") else CODE_TEXT


def slide_code(prs, s):
    slide = new_slide(prs)
    add_header(slide, s["num"], s["title"])
    y = Inches(1.4)
    if s.get("intro"):
        intro_h = Inches(0.42) * len(s["intro"]) + Inches(0.15)
        add_bullets(slide, MARGIN_X, y, SLIDE_W - 2 * MARGIN_X, intro_h, s["intro"], size=16, gap=4)
        y = y + intro_h
    code_h = Inches(6.0) - y - (Inches(0.75) if s.get("pregunta") else Emu(0))
    add_code_block(slide, MARGIN_X, y, SLIDE_W - 2 * MARGIN_X, code_h, s["code"])
    if s.get("pregunta"):
        add_pregunta(slide, s["pregunta"])
    add_footer(slide, s["num"])
    add_notes(slide, s.get("notes"))


def slide_twocol(prs, s):
    slide = new_slide(prs)
    add_header(slide, s["num"], s["title"])
    col_w = (SLIDE_W - 2 * MARGIN_X - Inches(0.4)) / 2
    y = Inches(1.5)
    # columna izquierda
    add_rect(slide, MARGIN_X, y, col_w, Inches(0.55), NAVY_LIGHT)
    add_textbox_simple(slide, MARGIN_X, y, col_w, Inches(0.55), s["left_title"], size=17, bold=True,
                        color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_bullets(slide, MARGIN_X, y + Inches(0.75), col_w, Inches(4.2), s["left_items"], size=16, gap=8)
    # columna derecha
    x2 = MARGIN_X + col_w + Inches(0.4)
    add_rect(slide, x2, y, col_w, Inches(0.55), TEAL)
    add_textbox_simple(slide, x2, y, col_w, Inches(0.55), s["right_title"], size=17, bold=True,
                        color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_bullets(slide, x2, y + Inches(0.75), col_w, Inches(4.2), s["right_items"], size=16, gap=8)
    add_footer(slide, s["num"])
    add_notes(slide, s.get("notes"))


RENDERERS = {
    "title": slide_title,
    "bullets": slide_bullets,
    "flow": slide_flow,
    "table": slide_table,
    "code": slide_code,
    "twocol": slide_twocol,
}
