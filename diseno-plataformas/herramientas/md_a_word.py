# -*- coding: utf-8 -*-
"""
Convierte un documento Markdown de la cátedra (por ejemplo, un trabajo práctico)
en un .docx con el encabezado institucional y el pie «Página X de Y».

Uso:
    python diseno-plataformas/herramientas/md_a_word.py diseno-plataformas/clases/trabajopractico-N.md
    python diseno-plataformas/herramientas/md_a_word.py diseno-plataformas/clases/trabajopractico-N.md -o salida.docx
    python diseno-plataformas/herramientas/md_a_word.py diseno-plataformas/clases/tp.md --encabezado "Otro encabezado"

Por defecto el .docx se genera junto al .md, con el mismo nombre.

Markdown soportado (el subconjunto que usan los documentos de diseno-plataformas/):
  - títulos #, ## y ###
  - párrafos; cada salto de línea simple se respeta como salto de línea
  - **negrita**, *cursiva* y `código` dentro de párrafos, listas y tablas
  - listas con «-» o «1.», anidadas por sangría
  - tablas con barras verticales (la primera fila es el encabezado)
  - «---»: el primero cierra la portada (todo lo anterior se centra);
    los siguientes se ignoran

Requiere: pip install python-docx
"""
import argparse
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ENCABEZADO_POR_DEFECTO = "Facultad de Ingeniería – Diseño según Plataforma"

FUENTE = "Calibri"
TAMANO_TEXTO = Pt(11)
TAMANO_TABLA = Pt(9.5)
COLOR_TITULOS = RGBColor(0x1F, 0x3A, 0x5F)
COLOR_CODIGO = RGBColor(0x8B, 0x1E, 0x3F)
FONDO_ENCABEZADO_TABLA = "D9E2F3"
SANGRIA_POR_NIVEL = Cm(0.75)

RE_TITULO = re.compile(r"^(#{1,3})\s+(.*)$")
RE_LISTA = re.compile(r"^(\s*)([-*]|\d+\.)\s+(.*)$")
RE_SEPARADOR_TABLA = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
RE_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<!\*)\*[^*\s][^*]*\*(?!\*))")


# ----------------------------------------------------------------------------
# Formato en línea
# ----------------------------------------------------------------------------
def agregar_texto(parrafo, texto, negrita=False, tamano=None):
    """Agrega texto con **negrita**, *cursiva* y `código` a un párrafo."""
    for parte in RE_INLINE.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
            run = parrafo.add_run(parte[2:-2])
            run.bold = True
        elif parte.startswith("`") and parte.endswith("`"):
            run = parrafo.add_run(parte[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = COLOR_CODIGO
            run.bold = negrita
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            run = parrafo.add_run(parte[1:-1])
            run.italic = True
            run.bold = negrita
        else:
            run = parrafo.add_run(parte)
            run.bold = negrita
        if tamano:
            run.font.size = tamano


# ----------------------------------------------------------------------------
# Encabezado y pie de página
# ----------------------------------------------------------------------------
def agregar_campo(parrafo, instruccion):
    """Inserta un campo de Word (PAGE, NUMPAGES) que se actualiza al abrir."""
    run = parrafo.add_run()
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruccion
    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")
    valor = OxmlElement("w:t")
    valor.text = "1"
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    for elemento in (inicio, instr, separador, valor, fin):
        run._r.append(elemento)
    return run


def configurar_pagina(documento, encabezado):
    seccion = documento.sections[0]
    seccion.page_height = Cm(29.7)
    seccion.page_width = Cm(21.0)
    seccion.top_margin = Cm(2.5)
    seccion.bottom_margin = Cm(2.2)
    seccion.left_margin = Cm(2.5)
    seccion.right_margin = Cm(2.0)

    p = seccion.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(encabezado)
    run.font.size = Pt(9)
    run.font.color.rgb = COLOR_TITULOS

    p = seccion.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for texto, campo in (("Página ", "PAGE"), (" de ", "NUMPAGES")):
        p.add_run(texto).font.size = Pt(9)
        agregar_campo(p, campo).font.size = Pt(9)


def configurar_estilos(documento):
    normal = documento.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = TAMANO_TEXTO
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    normal.paragraph_format.space_after = Pt(4)
    for nombre, tamano in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        estilo = documento.styles[nombre]
        estilo.font.name = FUENTE
        estilo.font.size = Pt(tamano)
        estilo.font.bold = True
        estilo.font.color.rgb = COLOR_TITULOS
        fuentes = estilo.element.rPr.rFonts
        for atributo in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fuentes.attrib.pop(qn(atributo), None)  # sin esto, Word usa la fuente del tema
        estilo.paragraph_format.space_before = Pt(14 if nombre != "Heading 3" else 10)
        estilo.paragraph_format.space_after = Pt(6)
        estilo.paragraph_format.keep_with_next = True


# ----------------------------------------------------------------------------
# Bloques
# ----------------------------------------------------------------------------
def sombrear_celda(celda, color):
    tc_pr = celda._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), color)
    tc_pr.append(sombra)


def repetir_fila_encabezado(fila):
    tr_pr = fila._tr.get_or_add_trPr()
    elemento = OxmlElement("w:tblHeader")
    elemento.set(qn("w:val"), "true")
    tr_pr.append(elemento)


def dividir_fila(linea):
    linea = linea.strip()
    if linea.startswith("|"):
        linea = linea[1:]
    if linea.endswith("|"):
        linea = linea[:-1]
    return [celda.strip() for celda in linea.split("|")]


def anchos_de_columna(filas, columnas, ancho_total):
    """Reparte el ancho según el texto más largo de cada columna (con mínimo y máximo)."""
    pesos = []
    for j in range(columnas):
        largo = max(len(f[j]) if j < len(f) else 0 for f in filas)
        pesos.append(min(max(largo, 8), 70))
    suma = sum(pesos)
    return [int(ancho_total * p / suma) for p in pesos]


def agregar_tabla(documento, lineas):
    filas = [dividir_fila(l) for l in lineas if not RE_SEPARADOR_TABLA.match(l.strip())]
    columnas = max(len(f) for f in filas)
    tabla = documento.add_table(rows=len(filas), cols=columnas)
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = False
    seccion = documento.sections[-1]
    anchos = anchos_de_columna(filas, columnas,
                               seccion.page_width - seccion.left_margin - seccion.right_margin)
    for j, columna in enumerate(tabla.columns):
        columna.width = anchos[j]
    for i, fila in enumerate(filas):
        for j in range(columnas):
            celda = tabla.cell(i, j)
            celda.width = anchos[j]  # Word toma el ancho de cada celda, no el de la columna
            parrafo = celda.paragraphs[0]
            parrafo.paragraph_format.space_after = Pt(0)
            texto = fila[j] if j < len(fila) else ""
            agregar_texto(parrafo, texto, negrita=(i == 0), tamano=TAMANO_TABLA)
            if i == 0:
                sombrear_celda(celda, FONDO_ENCABEZADO_TABLA)
        if i == 0:
            repetir_fila_encabezado(tabla.rows[0])
    documento.add_paragraph().paragraph_format.space_after = Pt(2)


def agregar_item_lista(documento, nivel, marcador, texto):
    parrafo = documento.add_paragraph()
    formato = parrafo.paragraph_format
    formato.left_indent = SANGRIA_POR_NIVEL * (nivel + 1)
    formato.first_line_indent = -Cm(0.6)
    formato.space_after = Pt(2)
    vineta = "•" if marcador in ("-", "*") else marcador
    if nivel > 0 and vineta == "•":
        vineta = "◦"
    parrafo.add_run(vineta + "\t")
    formato.tab_stops.add_tab_stop(SANGRIA_POR_NIVEL * (nivel + 1))
    agregar_texto(parrafo, texto)


# ----------------------------------------------------------------------------
# Conversión
# ----------------------------------------------------------------------------
def convertir(ruta_md, ruta_docx, encabezado):
    with open(ruta_md, encoding="utf-8") as archivo:
        lineas = archivo.read().splitlines()

    documento = Document()
    configurar_estilos(documento)
    configurar_pagina(documento, encabezado)

    en_portada = True
    parrafo_abierto = None
    sangrias_lista = []  # sangrías de los niveles de lista abiertos
    i = 0
    while i < len(lineas):
        linea = lineas[i]
        limpia = linea.strip()

        # línea vacía: cierra párrafos y listas
        if not limpia:
            parrafo_abierto = None
            sangrias_lista = []
            i += 1
            continue

        # separador horizontal
        if re.fullmatch(r"-{3,}|\*{3,}", limpia):
            parrafo_abierto = None
            en_portada = False
            i += 1
            continue

        # tabla: se toman todas las líneas consecutivas que empiezan con «|»
        if limpia.startswith("|"):
            bloque = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                bloque.append(lineas[i])
                i += 1
            agregar_tabla(documento, bloque)
            parrafo_abierto = None
            continue

        # títulos
        titulo = RE_TITULO.match(limpia)
        if titulo:
            nivel = len(titulo.group(1))
            estilo = "Title" if (nivel == 1 and en_portada) else f"Heading {nivel}"
            parrafo = documento.add_paragraph(style=estilo)
            agregar_texto(parrafo, titulo.group(2))
            if en_portada:
                parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            parrafo_abierto = None
            i += 1
            continue

        # listas
        item = RE_LISTA.match(linea)
        if item:
            sangria = len(item.group(1).expandtabs(4))
            while sangrias_lista and sangria < sangrias_lista[-1]:
                sangrias_lista.pop()
            if not sangrias_lista or sangria > sangrias_lista[-1]:
                sangrias_lista.append(sangria)
            agregar_item_lista(documento, len(sangrias_lista) - 1, item.group(2), item.group(3))
            parrafo_abierto = None
            i += 1
            continue

        # párrafo: las líneas consecutivas se unen con salto de línea
        if parrafo_abierto is None:
            parrafo_abierto = documento.add_paragraph()
            if en_portada:
                parrafo_abierto.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            parrafo_abierto.add_run().add_break()
        agregar_texto(parrafo_abierto, limpia)
        i += 1

    documento.save(ruta_docx)


def main():
    parser = argparse.ArgumentParser(description="Convierte un Markdown de la cátedra a Word (.docx).")
    parser.add_argument("markdown", help="ruta del archivo .md")
    parser.add_argument("-o", "--salida", help="ruta del .docx (por defecto, junto al .md)")
    parser.add_argument("--encabezado", default=ENCABEZADO_POR_DEFECTO,
                        help=f'texto del encabezado (por defecto: "{ENCABEZADO_POR_DEFECTO}")')
    args = parser.parse_args()

    if not os.path.isfile(args.markdown):
        sys.exit(f"No existe el archivo: {args.markdown}")
    salida = args.salida or os.path.splitext(args.markdown)[0] + ".docx"
    try:
        convertir(args.markdown, salida, args.encabezado)
    except PermissionError:
        sys.exit(f"No se pudo escribir {salida}: ¿está abierto en Word? Cerralo y volvé a ejecutar.")
    print(f"OK - {salida}")


if __name__ == "__main__":
    main()
