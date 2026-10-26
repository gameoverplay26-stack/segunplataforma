# -*- coding: utf-8 -*-
"""
Genera el deck de la Unidad 5 (Unidad-5-2026-completa.pptx) a partir de
diseno-plataformas/unidad-05/04-diapositivas.md, que es la unica fuente de contenido.

Reutiliza el generador de la Unidad 4 (misma paleta, layout, sintaxis y graficos nativos):
solo cambia la fuente, la carpeta de las imagenes, el destacado "PARA CONSOLA" y el nombre del archivo.

Uso (desde diseno-plataformas/unidad-05/):
    python scripts/graficos_consolas.py --slides   # graficos sin titulo, para las slides
    python scripts/build_decks.py                  # Unidad-5-2026-completa.pptx
    python scripts/build_decks.py --version 2      # Unidad-5-2026-completa-v2.pptx (no pisa la anterior)
Requiere: pip install python-pptx Pillow
"""
import argparse
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
UNIT_DIR = os.path.dirname(HERE)
U4_SCRIPTS = os.path.join(os.path.dirname(UNIT_DIR), "unidad-04", "scripts")
sys.path.insert(0, U4_SCRIPTS)

import build_decks as u4  # noqa: E402  (generador de la Unidad 4)
from pptx import Presentation  # noqa: E402

SOURCE = os.path.join(UNIT_DIR, "04-diapositivas.md")

# Las capturas y tarjetas se resuelven contra la carpeta de esta unidad.
u4.UNIT_DIR = UNIT_DIR
# Destacado propio de la unidad: "consola:" reemplaza a "movil:".
u4.KEYS = tuple(k for k in u4.KEYS if k != "movil") + ("consola",)
u4.CALLOUTS = (("consola", "PARA CONSOLA", u4.TEAL),) + tuple(c for c in u4.CALLOUTS if c[0] != "movil")


def footer(slide, deck, s, total):
    """Pie de la U4 con el numero de unidad de esta (el original dice 'Unidad IV').
    Desde la v2 la unidad se organiza en bloques que se dan segun el tiempo: el pie dice 'Bloque N'."""
    tf = u4.textbox(slide, u4.MARGIN, u4.FOOTER_Y, u4.SLIDE_W - 2 * u4.MARGIN, u4.Inches(0.3), u4.MSO_ANCHOR.MIDDLE)
    where = "Anexo" if deck["title"].startswith("Anexo") else "Bloque %d" % deck["num"]
    u4.run(tf.paragraphs[0], "Diseño según Plataformas de Juego · Unidad V · %s · %s" % (where, deck["title"]), 10, u4.MUTED)
    tf2 = u4.textbox(slide, u4.SLIDE_W - u4.MARGIN - u4.Inches(1), u4.FOOTER_Y, u4.Inches(1), u4.Inches(0.3), u4.MSO_ANCHOR.MIDDLE)
    p2 = tf2.paragraphs[0]
    p2.alignment = u4.PP_ALIGN.RIGHT
    u4.run(p2, "%d/%d" % (s["num"], total), 10, u4.MUTED)


u4.footer = footer


def build_unido(decks, out_dir, version=None):
    """Un solo .pptx con los bloques y los anexos, numeracion continua; el pie conserva la seccion."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = u4.SLIDE_W, u4.SLIDE_H
    total = sum(len(d["slides"]) for d in decks)
    n = 0
    for deck in decks:
        for s in deck["slides"]:
            n += 1
            s = dict(s, num=n)
            maker = u4.build_cover if s.get("tipo") == "portada" else u4.build_content
            u4.add_notes(maker(prs, deck, s, total), s["notes"])
    suffix = "-v%d" % version if version and version > 1 else ""
    path = os.path.join(out_dir, "Unidad-5-2026-completa%s.pptx" % suffix)
    prs.save(path)
    return path, total


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Genera el deck de la Unidad 5.")
    ap.add_argument("--version", type=int, help="agrega el sufijo -vN (N > 1) para no pisar la version anterior")
    ap.add_argument("--out-dir", default=UNIT_DIR, help="carpeta de salida (por defecto, la de la unidad)")
    args = ap.parse_args()
    decks = u4.parse(SOURCE)
    if not decks:
        sys.exit("No se encontraron secciones '# DECK' en %s" % SOURCE)
    path, total = build_unido(decks, args.out_dir, args.version)
    print("OK  %-40s %2d slides" % (os.path.basename(path), total))
